"""
Módulo para extraer texto de archivos DOCX y DOC.
"""
import os
import logging

log = logging.getLogger(__name__)

def extract_text_from_docx(file_path: str) -> tuple[str, bool]:
    """
    Extrae texto de un archivo DOCX o DOC.
    Retorna (texto, usado_ocr)
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.docx':
        return _extract_docx(file_path)
    elif ext == '.doc':
        return _extract_doc(file_path)
    else:
        return "", False

def _extract_docx(file_path: str) -> tuple[str, bool]:
    """Extrae texto de DOCX con python-docx."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = []

        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())

        # Extraer texto de tablas también
        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    paragraphs.append(row_text)

        text = '\n\n'.join(paragraphs)
        return text, False
    except Exception as e:
        log.error(f"Error extrayendo DOCX: {e}")
        return "", False

def _extract_doc(file_path: str) -> tuple[str, bool]:
    """Extrae texto de DOC usando antiword o catdoc."""
    try:
        import subprocess
        # Intentar con antiword
        result = subprocess.run(
            ['antiword', file_path],
            capture_output=True, text=True, timeout=30
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.strip(), False

        # Intentar con catdoc
        result = subprocess.run(
            ['catdoc', file_path],
            capture_output=True, text=True, timeout=30
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.strip(), False

        return "", False
    except Exception as e:
        log.error(f"Error extrayendo DOC: {e}")
        return "", False

def is_supported(filename: str) -> bool:
    """Verifica si el archivo es soportado."""
    ext = os.path.splitext(filename)[1].lower()
    return ext in ['.pdf', '.docx', '.doc']

def get_extension(filename: str) -> str:
    """Retorna la extensión del archivo."""
    return os.path.splitext(filename)[1].lower()
